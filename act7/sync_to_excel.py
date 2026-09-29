#!/usr/bin/env python3
"""
Sync Benchmark Results to 2110415 Storage Benchmark Template.xlsx
Usage: python sync_to_excel.py [--json-results benchmark_results.json] [--template "2110415 Storage Benchmark Template.xlsx"]
"""

import os
import sys
import json
import argparse
import openpyxl

ROW_MAPPINGS = {
    "2 KiB": {
        "start_row": 6,
        "iterations": [32768, 65536, 131072, 262144, 393216],
        "s3_iterations": [512, 1024, 2048, 4096, 6144]
    },
    "256 KiB": {
        "start_row": 17,
        "iterations": [512, 1024, 2048, 4096, 6144],
        "s3_iterations": [512, 1024, 2048, 4096, 6144]
    },
    "1 MiB": {
        "start_row": 28,
        "iterations": [128, 256, 512, 1024, 1536],
        "s3_iterations": [128, 256, 512, 1024, 1536]
    },
    "128 MiB": {
        "start_row": 39,
        "iterations": [1, 2, 4, 8, 12],
        "s3_iterations": [1, 2, 4, 8, 12]
    }
}

def sync_results(results_file: str, excel_file: str):
    if not os.path.exists(results_file):
        print(f"Error: Results file '{results_file}' not found.", file=sys.stderr)
        sys.exit(1)
    if not os.path.exists(excel_file):
        print(f"Error: Excel file '{excel_file}' not found.", file=sys.stderr)
        sys.exit(1)

    with open(results_file, "r", encoding="utf-8") as f:
        data = json.load(f)

    wb = openpyxl.load_workbook(excel_file)
    ws = wb.active

    for label, mapping in ROW_MAPPINGS.items():
        if label not in data:
            continue

        label_data = data[label]
        start_row = mapping["start_row"]

        # 1. Fill EBS & Ephemeral
        for idx, n in enumerate(mapping["iterations"]):
            row = start_row + idx
            str_n = str(n)

            # EBS: Cols B (Total Time), C (Per File), D (Bandwidth)
            if "ebs" in label_data and str_n in label_data["ebs"]:
                ebs_res = label_data["ebs"][str_n]
                ws[f"B{row}"] = round(ebs_res["total_write_time"], 4)
                ws[f"C{row}"] = round(ebs_res["write_time_per_file"], 6)
                ws[f"D{row}"] = round(ebs_res["bandwidth_mb_s"], 4)

            # Ephemeral: Cols E (Total Time), F (Per File), G (Bandwidth)
            if "ephemeral" in label_data and str_n in label_data["ephemeral"]:
                eph_res = label_data["ephemeral"][str_n]
                ws[f"E{row}"] = round(eph_res["total_write_time"], 4)
                ws[f"F{row}"] = round(eph_res["write_time_per_file"], 6)
                ws[f"G{row}"] = round(eph_res["bandwidth_mb_s"], 4)

        # 2. Fill S3
        for idx, n in enumerate(mapping["s3_iterations"]):
            row = start_row + idx
            str_n = str(n)

            # S3: Cols I (Total Time), J (Per File), K (Bandwidth)
            if "s3" in label_data and str_n in label_data["s3"]:
                s3_res = label_data["s3"][str_n]
                ws[f"I{row}"] = round(s3_res["total_write_time"], 4)
                ws[f"J{row}"] = round(s3_res["write_time_per_file"], 6)
                ws[f"K{row}"] = round(s3_res["bandwidth_mb_s"], 4)

    wb.save(excel_file)
    print(f"[+] Successfully populated benchmark data into '{excel_file}'.")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Sync benchmark JSON into Excel template")
    parser.add_argument("--json-results", default="benchmark_results.json", help="Path to benchmark_results.json")
    parser.add_argument("--template", default="2110415 Storage Benchmark Template.xlsx", help="Path to Excel template")
    args = parser.parse_args()

    sync_results(args.json_results, args.template)
