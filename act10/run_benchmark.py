import os
import sys
import time
import json
import base64
import re
import csv
import urllib.request
import boto3
from PIL import Image as PILImage
import openpyxl
from openpyxl.drawing.image import Image as XLImage

def parse_person_count(detections):
    if not detections:
        return 0
    m = re.search(r'(\d+)\s+person', detections)
    if m:
        return int(m.group(1))
    return 0

def main():
    print("=== Starting SDS Edge-Cloud People Counting Benchmark ===")
    
    # Setup directories
    img_dir = "sds-edge-computing/experiment_images"
    os.makedirs(img_dir, exist_ok=True)
    
    # 1. AWS SageMaker client setup
    with open('sds-sagemaker_accessKeys.csv', mode='r', encoding='utf-8-sig') as f:
        r = list(csv.DictReader(f))
        ak = r[0]['Access key ID']
        sk = r[0]['Secret access key']
    
    sm = boto3.client('sagemaker-runtime', region_name='ap-southeast-2', aws_access_key_id=ak, aws_secret_access_key=sk)
    
    # Ground truth mapping
    ground_truth = {
        1: 0,
        2: 9,
        3: 14,
        4: 24,
        5: 3,
        6: 2,
        7: 1,
        8: 3,
        9: 1,
        10: 5
    }
    
    # Load workbook
    excel_path = 'sds-edge-computing/People counting record sheet.xlsx'
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active
    
    # Set Col D width to match C, I, N
    ws.column_dimensions['D'].width = 25.1
    
    results = {}
    
    for img_idx in range(1, 11):
        raw_img_path = f'sds-edge-computing/{img_idx}.jpg'
        print(f"\n--- Processing Image {img_idx}/10: {raw_img_path} ---")
        
        with open(raw_img_path, 'rb') as f:
            raw_bytes = f.read()
        b64_img = base64.b64encode(raw_bytes).decode('utf-8')
        
        # Save resized original image for Excel
        orig_pil = PILImage.open(raw_img_path)
        orig_pil.thumbnail((180, 120))
        orig_thumb_path = f"{img_dir}/orig_{img_idx}.jpg"
        orig_pil.save(orig_thumb_path, "JPEG", quality=80)
        
        img_results = {
            'edge': {'resp_times': [], 'proc_times': [], 'count': 0, 'img_path': None},
            'cloud': {'resp_times': [], 'proc_times': [], 'count': 0, 'img_path': None},
            'edge_cloud': {'resp_times': [], 'proc_times': [], 'count': 0, 'img_path': None},
        }
        
        # --- 3 Trials for Edge ---
        print("  Running Edge (3 trials)...")
        for t in range(3):
            t0 = time.time()
            data = json.dumps({'image': b64_img}).encode('utf-8')
            req = urllib.request.Request('http://127.0.0.1:5001/predict', data=data, headers={'Content-Type': 'application/json'})
            with urllib.request.urlopen(req) as resp:
                res = json.loads(resp.read().decode('utf-8'))
            rt = (time.time() - t0) * 1000
            pt = res['processing_time'] * 1000
            cnt = parse_person_count(res.get('detections', ''))
            img_results['edge']['resp_times'].append(rt)
            img_results['edge']['proc_times'].append(pt)
            img_results['edge']['count'] = cnt
            if t == 0 and res.get('predicted_image'):
                pred_bytes = base64.b64decode(res['predicted_image'])
                p_path = f"{img_dir}/edge_{img_idx}.jpg"
                with open(p_path, 'wb') as pf:
                    pf.write(pred_bytes)
                p_pil = PILImage.open(p_path)
                p_pil.thumbnail((180, 120))
                p_pil.save(p_path, "JPEG", quality=80)
                img_results['edge']['img_path'] = p_path
            print(f"    Trial {t+1}: Resp={rt:.2f}ms, Proc={pt:.2f}ms, Count={cnt}")
        
        # --- 3 Trials for Cloud ---
        print("  Running Cloud (3 trials)...")
        for t in range(3):
            t0 = time.time()
            resp = sm.invoke_endpoint(EndpointName='YOLOv8', ContentType='application/json', Body=b64_img)
            rt = (time.time() - t0) * 1000
            res = json.loads(resp['Body'].read().decode('utf-8'))['data']
            pt = res['processing_time'] * 1000
            cnt = parse_person_count(res.get('detections', ''))
            img_results['cloud']['resp_times'].append(rt)
            img_results['cloud']['proc_times'].append(pt)
            img_results['cloud']['count'] = cnt
            if t == 0 and res.get('predicted_image'):
                pred_bytes = base64.b64decode(res['predicted_image'])
                p_path = f"{img_dir}/cloud_{img_idx}.jpg"
                with open(p_path, 'wb') as pf:
                    pf.write(pred_bytes)
                p_pil = PILImage.open(p_path)
                p_pil.thumbnail((180, 120))
                p_pil.save(p_path, "JPEG", quality=80)
                img_results['cloud']['img_path'] = p_path
            print(f"    Trial {t+1}: Resp={rt:.2f}ms, Proc={pt:.2f}ms, Count={cnt}")
            
        # --- 3 Trials for Edge-Cloud ---
        print("  Running Edge-Cloud (3 trials)...")
        for t in range(3):
            t0 = time.time()
            data = json.dumps({'image': b64_img}).encode('utf-8')
            req = urllib.request.Request('http://127.0.0.1:5002/predict', data=data, headers={'Content-Type': 'application/json'})
            with urllib.request.urlopen(req) as resp:
                res = json.loads(resp.read().decode('utf-8'))
            rt = (time.time() - t0) * 1000
            pt = res['processing_time'] * 1000
            cnt = parse_person_count(res.get('detections', ''))
            img_results['edge_cloud']['resp_times'].append(rt)
            img_results['edge_cloud']['proc_times'].append(pt)
            img_results['edge_cloud']['count'] = cnt
            if t == 0 and res.get('predicted_image'):
                pred_bytes = base64.b64decode(res['predicted_image'])
                p_path = f"{img_dir}/edge_cloud_{img_idx}.jpg"
                with open(p_path, 'wb') as pf:
                    pf.write(pred_bytes)
                p_pil = PILImage.open(p_path)
                p_pil.thumbnail((180, 120))
                p_pil.save(p_path, "JPEG", quality=80)
                img_results['edge_cloud']['img_path'] = p_path
            print(f"    Trial {t+1}: Resp={rt:.2f}ms, Proc={pt:.2f}ms, Count={cnt}, Model={res.get('used_model')}")
            
        results[img_idx] = img_results
        
        # --- Write to Excel ---
        base_row = 7 + (img_idx - 1) * 5
        r_resp = base_row + 1
        r_proc = base_row + 2
        r_data = base_row + 4
        
        # Ground Truth
        ws[f'C{r_data}'] = ground_truth[img_idx]
        
        # Edge values
        ws[f'E{r_resp}'] = round(img_results['edge']['resp_times'][0], 2)
        ws[f'F{r_resp}'] = round(img_results['edge']['resp_times'][1], 2)
        ws[f'G{r_resp}'] = round(img_results['edge']['resp_times'][2], 2)
        
        ws[f'E{r_proc}'] = round(img_results['edge']['proc_times'][0], 2)
        ws[f'F{r_proc}'] = round(img_results['edge']['proc_times'][1], 2)
        ws[f'G{r_proc}'] = round(img_results['edge']['proc_times'][2], 2)
        
        ws[f'D{r_data}'] = img_results['edge']['count']
        
        # Cloud values
        ws[f'J{r_resp}'] = round(img_results['cloud']['resp_times'][0], 2)
        ws[f'K{r_resp}'] = round(img_results['cloud']['resp_times'][1], 2)
        ws[f'L{r_resp}'] = round(img_results['cloud']['resp_times'][2], 2)
        
        ws[f'J{r_proc}'] = round(img_results['cloud']['proc_times'][0], 2)
        ws[f'K{r_proc}'] = round(img_results['cloud']['proc_times'][1], 2)
        ws[f'L{r_proc}'] = round(img_results['cloud']['proc_times'][2], 2)
        
        ws[f'I{r_data}'] = img_results['cloud']['count']
        
        # Edge-Cloud values
        ws[f'O{r_resp}'] = round(img_results['edge_cloud']['resp_times'][0], 2)
        ws[f'P{r_resp}'] = round(img_results['edge_cloud']['resp_times'][1], 2)
        ws[f'Q{r_resp}'] = round(img_results['edge_cloud']['resp_times'][2], 2)
        
        ws[f'O{r_proc}'] = round(img_results['edge_cloud']['proc_times'][0], 2)
        ws[f'P{r_proc}'] = round(img_results['edge_cloud']['proc_times'][1], 2)
        ws[f'Q{r_proc}'] = round(img_results['edge_cloud']['proc_times'][2], 2)
        
        ws[f'N{r_data}'] = img_results['edge_cloud']['count']
        
        # Clear placeholder text
        ws[f'C{base_row}'].value = None
        ws[f'D{base_row}'].value = None
        ws[f'I{base_row}'].value = None
        ws[f'N{base_row}'].value = None
        
        # Insert images
        xl_orig = XLImage(orig_thumb_path)
        xl_orig.width = 160
        xl_orig.height = 100
        ws.add_image(xl_orig, f'C{base_row}')
        
        if img_results['edge']['img_path']:
            xl_edge = XLImage(img_results['edge']['img_path'])
            xl_edge.width = 160
            xl_edge.height = 100
            ws.add_image(xl_edge, f'D{base_row}')
            
        if img_results['cloud']['img_path']:
            xl_cloud = XLImage(img_results['cloud']['img_path'])
            xl_cloud.width = 160
            xl_cloud.height = 100
            ws.add_image(xl_cloud, f'I{base_row}')
            
        if img_results['edge_cloud']['img_path']:
            xl_ec = XLImage(img_results['edge_cloud']['img_path'])
            xl_ec.width = 160
            xl_ec.height = 100
            ws.add_image(xl_ec, f'N{base_row}')
            
    # Save workbook
    wb.save(excel_path)
    file_size_mb = os.path.getsize(excel_path) / (1024 * 1024)
    print(f"\n=== Successfully saved {excel_path} (File size: {file_size_mb:.2f} MB) ===")
    
    # Save benchmark results to json for summary verification
    with open('sds-edge-computing/benchmark_summary.json', 'w') as jf:
        json.dump(results, jf, indent=2)

if __name__ == '__main__':
    main()
